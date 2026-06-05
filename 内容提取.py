import fitz
import os
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import sys
def set_run_font(run, font_name="宋体", font_size=11, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size)
    run.bold = bold
def rect_overlap_ratio(rect1, rect2):
    inter = rect1 & rect2
    if inter.is_empty:
        return 0
    area1 = rect1.get_area()
    if area1 == 0:
        return 0
    return inter.get_area() / area1
def get_yellow_rects(page, yellow_threshold=0.72):
    yellow_rects = []
    for annot in page.annots() or []:
        if annot.type[0] == 8:
            yellow_rects.append(fitz.Rect(annot.rect))
    for d in page.get_drawings():
        fill = d.get("fill")
        if not fill:
            continue
        r, g, b = fill[:3]
        if r > yellow_threshold and g > yellow_threshold and b < 0.75:
            yellow_rects.append(fitz.Rect(d["rect"]))
    return yellow_rects
def group_words_into_lines(words, page_width, y_tolerance=4):
    lines = []
    for word in words:
        placed = False
        for line in lines:
            if abs(line["y"] - word["y0"]) <= y_tolerance:
                line["words"].append(word)
                line["ys"].append(word["y0"])
                placed = True
                break
        if not placed:
            lines.append({
                "y": word["y0"],
                "ys": [word["y0"]],
                "words": [word]
            })
    result_lines = []
    for line in lines:
        line_words = sorted(line["words"], key=lambda x: x["x0"])
        text_parts = []
        previous_x1 = None
        for w in line_words:
            if previous_x1 is None:
                text_parts.append(w["text"])
            else:
                gap = w["x0"] - previous_x1
                if gap > 20:
                    text_parts.append("    " + w["text"])
                else:
                    text_parts.append(" " + w["text"])
            previous_x1 = w["x1"]
        text = "".join(text_parts).strip()
        result_lines.append({
            "text": text,
            "y": sum(line["ys"]) / len(line["ys"]),
            "y1": max(w["y1"] for w in line_words),
            "x0": min(w["x0"] for w in line_words),
            "x1": max(w["x1"] for w in line_words),
            "page_width": page_width,
            "words": line_words
        })
    result_lines.sort(key=lambda x: x["y"])
    return result_lines
def extract_yellow_text(pdf_path, yellow_threshold=0.9, overlap_threshold=0.1, expand=0):
    doc = fitz.open(pdf_path)
    all_lines = []
    for page in doc:
        words = page.get_text("words")
        yellow_rects = get_yellow_rects(page, yellow_threshold)
        if not yellow_rects:
            continue
        selected_words = []
        for w in words:
            word_rect = fitz.Rect(w[:4])
            center = fitz.Point(
                (word_rect.x0 + word_rect.x1) / 2,
                word_rect.y0 + (word_rect.y1 - word_rect.y0) * 0.6
            )
            for yellow_rect in yellow_rects:
                rect = fitz.Rect(yellow_rect)
                rect.x0 -= expand
                rect.y0 -= expand
                rect.x1 += expand
                rect.y1 += expand
                # 参考之前 Streamlit 版本：只提取中心点在黄色区域内的文字
                if rect.contains(center):
                    selected_words.append({
                        "x0": w[0],
                        "y0": w[1],
                        "x1": w[2],
                        "y1": w[3],
                        "text": w[4],
                    })
                    break
        if not selected_words:
            continue
        unique = {}
        for w in selected_words:
            key = (
                round(w["x0"], 1),
                round(w["y0"], 1),
                w["text"]
            )
            unique[key] = w
        selected_words = list(unique.values())
        selected_words.sort(key=lambda x: (x["y0"], x["x0"]))
        lines = group_words_into_lines(
            selected_words,
            page_width=page.rect.width
        )
        for line in lines:
            line["page_index"] = page.number
        all_lines.extend(lines)
    doc.close()
    return all_lines
def is_table_like_line(line):
    words = line.get("words", [])
    if len(words) < 2:
        return False
    words = sorted(words, key=lambda w: w["x0"])
    big_gap_count = 0
    for i in range(1, len(words)):
        gap = words[i]["x0"] - words[i - 1]["x1"]
        if gap > 8:
            big_gap_count += 1
    return big_gap_count >= 1

def get_default_margins(lines):
    left_keys = []
    right_keys = []

    for line in lines:
        if not line["text"].strip():
            continue

        left_keys.append(round(line["x0"] / 10) * 10)
        right_keys.append(round((line["page_width"] - line["x1"]) / 10) * 10)

    default_left = max(set(left_keys), key=left_keys.count) if left_keys else 0
    default_right = max(set(right_keys), key=right_keys.count) if right_keys else 0

    return default_left, default_right


def is_layout_special_line(line, default_left, default_right, gap_threshold=18):
    words = line.get("words", [])

    if len(words) < 2:
        return False

    words = sorted(words, key=lambda w: w["x0"])

    left_blank = round(line["x0"] / 10) * 10
    right_blank = round((line["page_width"] - line["x1"]) / 10) * 10

    left_normal = abs(left_blank - default_left) <= 30
    right_normal = abs(right_blank - default_right) <= 30

    left_abnormal = not left_normal
    right_abnormal = not right_normal

    has_left_content = False
    has_right_content = False

    for i in range(1, len(words)):
        gap = words[i]["x0"] - words[i - 1]["x1"]

        if gap > gap_threshold:
            has_left_content = True
            has_right_content = True
            break

    if not (has_left_content and has_right_content):
        return False

    if left_normal and right_abnormal:
        return True

    if left_abnormal and right_normal:
        return True

    if left_abnormal and right_abnormal:
        return True

    return False
def get_table_areas(page, min_length=25, merge_gap=20):
    lines = []

    for d in page.get_drawings():
        for item in d.get("items", []):
            if item[0] != "l":
                continue
            p1 = item[1]
            p2 = item[2]
            x0, y0 = p1.x, p1.y
            x1, y1 = p2.x, p2.y
            length = ((x1 - x0) ** 2 + (y1 - y0) ** 2) ** 0.5
            if length < min_length:
                continue
            is_horizontal = abs(y1 - y0) <= 2
            is_vertical = abs(x1 - x0) <= 2
            if not (is_horizontal or is_vertical):
                continue
            rect = fitz.Rect(
                min(x0, x1) - 2,
                min(y0, y1) - 2,
                max(x0, x1) + 2,
                max(y0, y1) + 2
            )
            lines.append({
                "rect": rect,
                "type": "h" if is_horizontal else "v"
            })
    areas = []
    for line in lines:
        rect = line["rect"]
        added = False
        for area in areas:
            expanded = fitz.Rect(
                area["rect"].x0 - merge_gap,
                area["rect"].y0 - merge_gap,
                area["rect"].x1 + merge_gap,
                area["rect"].y1 + merge_gap
            )
            if expanded.intersects(rect):
                area["rect"].x0 = min(area["rect"].x0, rect.x0)
                area["rect"].y0 = min(area["rect"].y0, rect.y0)
                area["rect"].x1 = max(area["rect"].x1, rect.x1)
                area["rect"].y1 = max(area["rect"].y1, rect.y1)
                area["types"].append(line["type"])
                added = True
                break
        if not added:
            areas.append({
                "rect": fitz.Rect(rect),
                "types": [line["type"]]
            })
    table_areas = []
    for area in areas:
        h_count = area["types"].count("h")
        v_count = area["types"].count("v")
        if h_count >= 2 and v_count >= 2:
            r = area["rect"]
            table_areas.append(
                fitz.Rect(
                    max(0, r.x0 - 10),
                    max(0, r.y0 - 10),
                    min(page.rect.width, r.x1 + 10),
                    min(page.rect.height, r.y1 + 10)
                )
            )

    return table_areas
def get_table_area_for_line(line, table_areas):
    line_rect = fitz.Rect(
        line["x0"],
        line["y"],
        line["x1"],
        line.get("y1", line["y"] + 5)
    )
    for area in table_areas:
        if area.intersects(line_rect):
            return area
    return None
def merge_lines(lines):
    paragraphs = []
    current_paragraph = ""
    right_blank_keys = []
    left_blanks = []
    for line in lines:
        text = line["text"].strip()
        if not text:
            continue
        right_blank = line["page_width"] - line["x1"]
        left_blank = line["x0"]
        right_blank_key = round(right_blank / 10) * 10
        right_blank_keys.append(right_blank_key)
        left_blanks.append(left_blank)
    if right_blank_keys:
        default_right_blank = max(
            set(right_blank_keys),
            key=right_blank_keys.count
        )
    else:
        default_right_blank = 0
    default_left_blank = min(left_blanks) if left_blanks else 0
    for line in lines:
        text = line["text"].strip()
        if not text:
            continue
        right_blank = line["page_width"] - line["x1"]
        left_blank = line["x0"]
        right_key = round(right_blank / 10) * 10
        right_equal = right_key == default_right_blank
        left_equal = abs(left_blank - default_left_blank) <= 20
        if right_equal and left_equal:
            if current_paragraph:
                current_paragraph += text
            else:
                current_paragraph = text
            continue
        if not right_equal:
            if current_paragraph:
                current_paragraph += text
            else:
                current_paragraph = text
            paragraphs.append(current_paragraph)
            current_paragraph = ""
            continue
        if right_equal and not left_equal:
            if current_paragraph:
                paragraphs.append(current_paragraph)
            current_paragraph = text
            continue
    if current_paragraph:
        paragraphs.append(current_paragraph)
    return paragraphs
def write_paragraphs_to_docx(docx_file, text_buffer):
    if not text_buffer:
        return
    paragraphs = merge_lines(text_buffer)
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        p = docx_file.add_paragraph()
        run = p.add_run(para)
        set_run_font(run, "宋体", 11, False)
def make_clip_rect(page, table_lines):
    x0 = min(line["x0"] for line in table_lines)
    x1 = max(line["x1"] for line in table_lines)
    y0 = min(line["y"] for line in table_lines)
    y1 = max(line.get("y1", line["y"]) for line in table_lines)
    return fitz.Rect(
        max(0, x0 - 20),
        max(0, y0 - 8),
        min(page.rect.width, x1 + 20),
        min(page.rect.height, y1 + 7)
    )
def insert_table_image(word_doc, pdf_path, page_index, clip, zoom=3):
    doc = fitz.open(pdf_path)
    page = doc[page_index]
    pix = page.get_pixmap(
        matrix=fitz.Matrix(zoom, zoom),
        clip=clip,
        alpha=False
    )
    img_stream = BytesIO(pix.tobytes("png"))
    img_stream.seek(0)
    p = word_doc.add_paragraph()
    run = p.add_run()
    run.add_picture(img_stream, width=Inches(5.6))
    doc.close()
def line_in_block(line, block):
    if line["page_index"] != block["page_index"]:
        return False
    line_rect = fitz.Rect(
        line["x0"],
        line["y"],
        line["x1"],
        line.get("y1", line["y"] + 5)
    )
    return rect_overlap_ratio(line_rect, block["clip"]) > 0.02
def save_docx_in_order(lines, output_path, pdf_path):
    docx_file = Document()

    style = docx_file.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)

    pdf_doc = fitz.open(pdf_path)

    page_table_areas = {}
    for page in pdf_doc:
        page_table_areas[page.number] = get_table_areas(page)

    default_left, default_right = get_default_margins(lines)

    blocks = []
    used_table_clips = []

    i = 0

    while i < len(lines):
        line = lines[i]
        page_index = line["page_index"]

        table_area = get_table_area_for_line(
            line,
            page_table_areas.get(page_index, [])
        )

        if table_area:
            duplicated = False

            for used in used_table_clips:
                if used["page_index"] == page_index and used["clip"].intersects(table_area):
                    duplicated = True
                    break

            if not duplicated:
                blocks.append({
                    "page_index": page_index,
                    "start_index": i,
                    "end_index": i + 1,
                    "clip": table_area
                })

                used_table_clips.append({
                    "page_index": page_index,
                    "clip": table_area
                })
            i += 1
            continue
        if is_table_like_line(line) or is_layout_special_line(line, default_left, default_right):
            table_lines = [line]
            last_y = line["y"]
            j = i + 1
            while j < len(lines):
                next_line = lines[j]
                if next_line.get("page_index") != page_index:
                    break
                y_close = abs(next_line["y"] - last_y) <= 28
                next_table_area = get_table_area_for_line(
                    next_line,
                    page_table_areas.get(page_index, [])
                )
                same_kind = (
                    next_table_area
                    or is_table_like_line(next_line)
                    or is_layout_special_line(next_line, default_left, default_right)
                )
                if y_close and same_kind:
                    table_lines.append(next_line)
                    last_y = next_line["y"]
                    j += 1
                else:
                    break
            page = pdf_doc[page_index]
            clip = make_clip_rect(page, table_lines)
            blocks.append({
                "page_index": page_index,
                "start_index": i,
                "end_index": j,
                "clip": clip
            })
            i = j
        else:
            i += 1
    merged_blocks = []
    for block in blocks:
        if not merged_blocks:
            merged_blocks.append(block)
            continue
        last = merged_blocks[-1]
        same_page = block["page_index"] == last["page_index"]
        vertical_gap = block["clip"].y0 - last["clip"].y1
        if same_page and vertical_gap <= 25:
            last["end_index"] = max(last["end_index"], block["end_index"])
            last["clip"] = fitz.Rect(
                min(last["clip"].x0, block["clip"].x0),
                min(last["clip"].y0, block["clip"].y0),
                max(last["clip"].x1, block["clip"].x1),
                max(last["clip"].y1, block["clip"].y1)
            )
        else:
            merged_blocks.append(block)
    blocks = merged_blocks
    pdf_doc.close()
    text_buffer = []
    i = 0
    while i < len(lines):
        line = lines[i]
        block_start = None
        for block in blocks:
            if block["start_index"] == i:
                block_start = block
                break
        if block_start:
            write_paragraphs_to_docx(docx_file, text_buffer)
            text_buffer = []
            insert_table_image(
                docx_file,
                pdf_path,
                block_start["page_index"],
                block_start["clip"]
            )
            i = block_start["end_index"]
            continue
        if any(line_in_block(line, block) for block in blocks):
            i += 1
            continue
        text_buffer.append(line)
        i += 1
    write_paragraphs_to_docx(docx_file, text_buffer)
    docx_file.save(output_path)
def get_current_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))
def main():
    current_dir = get_current_dir()
    print("程序目录：", current_dir)
    all_files = os.listdir(current_dir)
    print("目录中的全部文件：")
    for f in all_files:
        print(repr(f))
    pdf_files = [
        f for f in os.listdir(current_dir)
        if f.lower().endswith(".pdf")
    ]
    print("识别到的 PDF：", pdf_files)
    if not pdf_files:
        input("当前文件夹中没有找到 PDF 文件。按回车退出...")
        return
    for pdf_file in pdf_files:
        pdf_path = os.path.join(current_dir, pdf_file)
        name_without_ext = os.path.splitext(pdf_file)[0]
        output_path = os.path.join(
            current_dir,
            f"{name_without_ext}_重点提取.docx"
        )
        print(f"正在处理：{pdf_file}")
        lines = extract_yellow_text(pdf_path)
        if not lines:
            print(f"未提取到黄色标注内容：{pdf_file}")
            continue
        save_docx_in_order(
            lines,
            output_path,
            pdf_path
        )
        print(f"已生成：{output_path}")
    input("全部处理完成，按回车退出...")
if __name__ == "__main__":
    main()
